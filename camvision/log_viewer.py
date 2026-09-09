"""Password-protected, read-only viewer for CamVision audit databases."""

from __future__ import annotations

import argparse
import base64
import hashlib
import hmac
import json
import os
import sqlite3
import sys
from pathlib import Path

from .audit import audit_root
from .config import ConfigManager


def _security_file(root: Path) -> Path:
    return root / "security.json"


def _hash(password: str, salt: bytes) -> str:
    return base64.b64encode(hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1)).decode("ascii")


def password_is_set(root: Path) -> bool:
    return _security_file(root).is_file()


def set_password(root: Path, password: str) -> None:
    if len(password) < 8:
        raise ValueError("Use at least 8 characters.")
    root.mkdir(parents=True, exist_ok=True)
    salt = os.urandom(16)
    target = _security_file(root)
    target.write_text(json.dumps({"salt": base64.b64encode(salt).decode("ascii"), "hash": _hash(password, salt)}), encoding="utf-8")
    try:
        os.chmod(root, 0o700)
        os.chmod(target, 0o600)
    except OSError:
        pass


def verify_password(root: Path, password: str) -> bool:
    try:
        saved = json.loads(_security_file(root).read_text(encoding="utf-8"))
        salt = base64.b64decode(saved["salt"])
        return hmac.compare_digest(saved["hash"], _hash(password, salt))
    except (OSError, KeyError, ValueError, TypeError):
        return False


def database_paths(root: Path) -> list[Path]:
    return sorted(root.glob("*/*/*.sqlite3"), reverse=True)


def read_events(root: Path) -> list[tuple]:
    rows: list[tuple] = []
    for path in database_paths(root):
        try:
            with sqlite3.connect(str(path)) as conn:
                rows.extend(conn.execute("SELECT local_time, severity, category, action, operator, program_name, message, details_json, snapshot_json FROM events ORDER BY id").fetchall())
        except sqlite3.Error:
            continue
    return sorted(rows, key=lambda row: row[0], reverse=True)


def _export_xlsx(rows: list[tuple], target: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit log"
    headers = ["Local time", "Severity", "Category", "Action", "Operator", "Program", "Message", "Details", "Machine snapshot"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    colors = {"error": "FFC7CE", "warning": "FFEB9C", "info": "C6EFCE"}
    for row in rows:
        ws.append(row)
        fill = colors.get(str(row[1]).lower())
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in {"A": 27, "B": 12, "C": 14, "D": 16, "E": 20, "F": 24, "G": 55, "H": 60, "I": 60}.items():
        ws.column_dimensions[column].width = width
    wb.save(target)


def main(argv=None) -> int:
    from PyQt5.QtGui import QColor
    from PyQt5.QtWidgets import (QApplication, QFileDialog, QHBoxLayout, QInputDialog,
                                 QLabel, QLineEdit, QMessageBox, QPushButton,
                                 QPlainTextEdit, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget)

    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True)
    args = parser.parse_args(argv)
    config = ConfigManager(args.config)
    root = audit_root()
    app = QApplication.instance() or QApplication(sys.argv)

    if password_is_set(root):
        password, ok = QInputDialog.getText(None, "Audit log", "Password:", QLineEdit.Password)
        if not ok or not verify_password(root, password):
            QMessageBox.warning(None, "Audit log", "Incorrect password.")
            return 1
    else:
        password, ok = QInputDialog.getText(None, "Set audit password", "New password (minimum 8 characters):", QLineEdit.Password)
        if not ok:
            return 0
        try:
            set_password(root, password)
        except ValueError as exc:
            QMessageBox.warning(None, "Audit log", str(exc))
            return 1

    window = QWidget()
    window.setWindowTitle("CamVision Audit Log Viewer")
    layout = QVBoxLayout(window)
    title = QLabel("Audit log — read-only")
    title.setStyleSheet("font-weight:bold;font-size:15px")
    layout.addWidget(title)
    filter_edit = QLineEdit()
    filter_edit.setPlaceholderText("Filter operator, program, event, or message")
    layout.addWidget(filter_edit)
    table = QTableWidget()
    headers = ["Time", "Level", "Category", "Action", "Operator", "Program", "Message"]
    table.setColumnCount(len(headers)); table.setHorizontalHeaderLabels(headers)
    table.setEditTriggers(QTableWidget.NoEditTriggers)
    table.setSelectionBehavior(QTableWidget.SelectRows)
    layout.addWidget(table)
    details = QPlainTextEdit()
    details.setReadOnly(True)
    details.setPlaceholderText("Select an event to view G54/G55, machine Z, and program safety values.")
    details.setMaximumHeight(155)
    layout.addWidget(details)
    buttons = QHBoxLayout()
    export = QPushButton("Export to Excel")
    reset = QPushButton("Reset password")
    reset.setVisible(config.checkbox("development_mode", False))
    buttons.addWidget(export); buttons.addWidget(reset); buttons.addStretch()
    layout.addLayout(buttons)
    rows = read_events(root)
    shown_rows = rows

    def populate() -> None:
        nonlocal shown_rows
        needle = filter_edit.text().casefold().strip()
        shown_rows = [r for r in rows if not needle or needle in " ".join(str(v) for v in r[:7]).casefold()]
        table.setRowCount(len(shown_rows))
        shades = {"error": "#ffc7ce", "warning": "#ffeb9c", "info": "#c6efce"}
        for row_no, row in enumerate(shown_rows):
            for col, value in enumerate(row[:7]):
                item = QTableWidgetItem(str(value or ""))
                item.setBackground(QColor(shades.get(str(row[1]).lower(), "#d9eaf7")))
                table.setItem(row_no, col, item)
        table.resizeColumnsToContents()

    def do_export() -> None:
        target, _ = QFileDialog.getSaveFileName(window, "Export audit log", "camvision-audit.xlsx", "Excel (*.xlsx)")
        if target:
            if not target.lower().endswith(".xlsx"):
                target += ".xlsx"
            try:
                _export_xlsx(shown_rows, target)
                QMessageBox.information(window, "Audit log", f"Exported {len(shown_rows)} records.")
            except Exception as exc:
                QMessageBox.warning(window, "Audit log", f"Could not export Excel file: {exc}")

    def show_details() -> None:
        selected = table.currentRow()
        if selected < 0 or selected >= len(shown_rows):
            return
        row = shown_rows[selected]
        try:
            detail_json = json.dumps(json.loads(row[7]), indent=2, sort_keys=True)
            snapshot_json = json.dumps(json.loads(row[8]), indent=2, sort_keys=True)
        except (TypeError, ValueError):
            detail_json, snapshot_json = str(row[7]), str(row[8])
        details.setPlainText(f"Program / event values:\n{detail_json}\n\nLinuxCNC snapshot:\n{snapshot_json}")

    def do_reset() -> None:
        password, ok = QInputDialog.getText(window, "Reset audit password", "New password:", QLineEdit.Password)
        if not ok:
            return
        try:
            set_password(root, password)
        except ValueError as exc:
            QMessageBox.warning(window, "Audit log", str(exc))
            return
        from .audit import AuditLog
        audit = AuditLog(root)
        audit.record("security", "password_reset", severity="warning",
                     message="Audit password reset while Development Mode was enabled.")
        audit.close()
        QMessageBox.information(window, "Audit log", "Password reset and recorded.")

    filter_edit.textChanged.connect(populate); export.clicked.connect(do_export); reset.clicked.connect(do_reset)
    table.itemSelectionChanged.connect(show_details)
    populate(); window.resize(1150, 650); window.show()
    return app.exec_()


if __name__ == "__main__":
    raise SystemExit(main())
